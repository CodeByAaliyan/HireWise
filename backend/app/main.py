from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Depends, Header
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, EmailStr, Field
import fitz
from docx import Document
from io import BytesIO
import hashlib
import hmac
import secrets
import sqlite3
import os
from pathlib import Path
from datetime import datetime, timedelta, timezone

from .analyzer import analyze_resume
from .llm import enhanced_feedback, generate_questions

app = FastAPI(title="HireWise API", version="1.0.0",
              description="Resume and job-description matching assistant")

DB_PATH = Path(__file__).resolve().parent.parent / "hirewise.db"
SESSION_DAYS = 7


def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    with get_db() as conn:
        conn.execute("""CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            email TEXT NOT NULL UNIQUE,
            password_hash TEXT NOT NULL,
            created_at TEXT NOT NULL
        )""")
        conn.execute("""CREATE TABLE IF NOT EXISTS sessions (
            token_hash TEXT PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
            expires_at TEXT NOT NULL
        )""")


init_db()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://127.0.0.1:5173"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

MAX_FILE_SIZE = 5 * 1024 * 1024

async def extract_text(upload: UploadFile) -> str:
    raw = await upload.read()
    if len(raw) > MAX_FILE_SIZE:
        raise HTTPException(status_code=413, detail="File is too large. Maximum size is 5 MB.")
    name = (upload.filename or "").lower()
    try:
        if name.endswith(".pdf"):
            with fitz.open(stream=raw, filetype="pdf") as pdf:
                return "\n".join(page.get_text() for page in pdf)
        if name.endswith(".docx"):
            doc = Document(BytesIO(raw))
            chunks = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    chunks.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(chunks)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Could not read the resume file: {exc}")
    raise HTTPException(status_code=400, detail="Unsupported file. Upload a PDF or DOCX resume.")



class AuthPayload(BaseModel):
    name: str = Field(default="", max_length=80)
    email: EmailStr
    password: str = Field(min_length=8, max_length=128)


def hash_password(password: str, salt: bytes | None = None) -> str:
    salt = salt or secrets.token_bytes(16)
    derived = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, 310_000)
    return f"{salt.hex()}:{derived.hex()}"


def verify_password(password: str, stored: str) -> bool:
    try:
        salt_hex, digest_hex = stored.split(":", 1)
        expected = bytes.fromhex(digest_hex)
        actual = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), bytes.fromhex(salt_hex), 310_000)
        return hmac.compare_digest(actual, expected)
    except (ValueError, TypeError):
        return False


def public_user(row):
    return {"id": row["id"], "name": row["name"], "email": row["email"]}


def current_user(authorization: str | None = Header(default=None)):
    if not authorization or not authorization.lower().startswith("bearer "):
        raise HTTPException(status_code=401, detail="Please log in to continue.")
    token = authorization.split(" ", 1)[1].strip()
    token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
    now = datetime.now(timezone.utc).isoformat()
    with get_db() as conn:
        row = conn.execute("""SELECT users.id, users.name, users.email
            FROM sessions JOIN users ON users.id = sessions.user_id
            WHERE sessions.token_hash = ? AND sessions.expires_at > ?""", (token_hash, now)).fetchone()
    if not row:
        raise HTTPException(status_code=401, detail="Your session has expired. Please log in again.")
    return public_user(row)


def create_session(user_id: int):
    token = secrets.token_urlsafe(40)
    token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
    expires = (datetime.now(timezone.utc) + timedelta(days=SESSION_DAYS)).isoformat()
    with get_db() as conn:
        conn.execute("INSERT INTO sessions(token_hash, user_id, expires_at) VALUES (?, ?, ?)", (token_hash, user_id, expires))
    return token


@app.post("/api/auth/register")
def register(body: AuthPayload):
    name = body.name.strip()
    email = str(body.email).lower().strip()
    if not name:
        raise HTTPException(status_code=400, detail="Please enter your name.")
    if len(body.password) < 8:
        raise HTTPException(status_code=400, detail="Password must be at least 8 characters.")
    try:
        with get_db() as conn:
            cur = conn.execute("INSERT INTO users(name, email, password_hash, created_at) VALUES (?, ?, ?, ?)",
                (name, email, hash_password(body.password), datetime.now(timezone.utc).isoformat()))
            user_id = cur.lastrowid
    except sqlite3.IntegrityError:
        raise HTTPException(status_code=409, detail="An account with this email already exists.")
    token = create_session(user_id)
    return {"token": token, "user": {"id": user_id, "name": name, "email": email}}


@app.post("/api/auth/login")
def login(body: AuthPayload):
    email = str(body.email).lower().strip()
    with get_db() as conn:
        row = conn.execute("SELECT * FROM users WHERE email = ?", (email,)).fetchone()
    if not row or not verify_password(body.password, row["password_hash"]):
        raise HTTPException(status_code=401, detail="Email or password is incorrect.")
    token = create_session(row["id"])
    return {"token": token, "user": public_user(row)}


@app.get("/api/auth/me")
def get_me(user=Depends(current_user)):
    return {"user": user}


@app.post("/api/auth/logout")
def logout(authorization: str | None = Header(default=None)):
    if authorization and authorization.lower().startswith("bearer "):
        token = authorization.split(" ", 1)[1].strip()
        token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
        with get_db() as conn:
            conn.execute("DELETE FROM sessions WHERE token_hash = ?", (token_hash,))
    return {"message": "Logged out successfully."}


@app.get("/health")
def health():
    return {"status": "ok", "service": "HireWise API"}

@app.post("/api/analyze")
async def analyze(
    resume: UploadFile = File(...),
    job_description: str = Form(...),
    use_ai: bool = Form(False),
    user=Depends(current_user),
):
    if not job_description.strip():
        raise HTTPException(status_code=400, detail="Please provide a job description.")
    resume_text = await extract_text(resume)
    if len(resume_text.strip()) < 30:
        raise HTTPException(status_code=400, detail="Very little text was extracted. If this is a scanned PDF, use an OCR-readable PDF or DOCX.")
    result = analyze_resume(resume_text, job_description)
    if use_ai:
        result["ai_feedback"] = enhanced_feedback(resume_text, job_description, result)
        result["ai_enabled"] = bool(result["ai_feedback"])
    else:
        result["ai_feedback"] = None
        result["ai_enabled"] = False
    return result

class QuestionsRequest(BaseModel):
    job_description: str
    resume_text: str = ""

@app.post("/api/interview-questions")
def interview_questions(body: QuestionsRequest, user=Depends(current_user)):
    if not body.job_description.strip():
        raise HTTPException(status_code=400, detail="Please provide a job description.")
    questions = generate_questions(body.job_description, body.resume_text)
    if not questions:
        questions = [
            "Walk me through a project relevant to this role and explain your contribution.",
            "Which technical skill in the job description are you strongest in, and how have you applied it?",
            "Describe a difficult problem you solved and how you evaluated the result.",
            "How would you approach a task when requirements are incomplete?",
            "Tell me about a time you received feedback and acted on it.",
            "How do you test your work before handing it over?",
            "Explain a trade-off you made in a project and what you learned.",
            "What would you want to learn first if you joined this team?"
        ]
        source = "starter questions (configure Gemini for tailored generation)"
    else:
        source = "Gemini-generated questions"
    return {"questions": questions, "source": source}
